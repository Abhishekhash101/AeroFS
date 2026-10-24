#!/usr/bin/env python3
"""
AeroFS Markdown to Word (.docx) Generator
Generates an executive, highly-styled Word document from README.md
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = "/home/abhishek/Documents/os_project"
README_PATH = os.path.join(WORKSPACE_DIR, "README.md")
OUTPUT_DOCX = os.path.join(WORKSPACE_DIR, "AeroFS_Project_Report.docx")
OUTPUT_DOCX_ALT = os.path.join(WORKSPACE_DIR, "AeroFS_Documentation.docx")

# Color Palette (Deep Navy / Slate Executive)
NAVY_HEX = "1E3A8A"
NAVY_RGB = RGBColor(30, 58, 138)
BLUE_HEX = "2563EB"
BLUE_RGB = RGBColor(37, 99, 235)
DARK_TEXT_RGB = RGBColor(31, 41, 55)    # #1F2937
MUTED_TEXT_RGB = RGBColor(100, 116, 139) # #64748B
BG_ALT_ROW_HEX = "F8FAFC"
BG_CODE_HEX = "F1F5F9"
BORDER_HEX = "CBD5E1"

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets internal padding (in dxa) of a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Sets clean thin borders around and inside a table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def parse_inline_tokens(text):
    """Parses markdown inline tokens: **bold**, *italic*, `code`, and [link](url)."""
    # Pattern to match bold, italic, code, link
    pattern = re.compile(
        r'(\*\*(.*?)\*\*)|(\*(.*?)\*)|(`(.*?)`)|(\[(.*?)\]\((.*?)\))'
    )
    tokens = []
    pos = 0
    for m in pattern.finditer(text):
        start, end = m.span()
        if start > pos:
            tokens.append(('text', text[pos:start]))
        if m.group(1): # bold
            tokens.append(('bold', m.group(2)))
        elif m.group(3): # italic
            tokens.append(('italic', m.group(4)))
        elif m.group(5): # code
            tokens.append(('code', m.group(6)))
        elif m.group(7): # link
            tokens.append(('link', m.group(8), m.group(9)))
        pos = end
    if pos < len(text):
        tokens.append(('text', text[pos:]))
    return tokens

def add_inline_runs(paragraph, text, base_font="Calibri", base_size=Pt(10.5), base_color=DARK_TEXT_RGB):
    """Applies inline formatting to a paragraph."""
    tokens = parse_inline_tokens(text)
    for token in tokens:
        ttype = token[0]
        if ttype == 'text':
            run = paragraph.add_run(token[1])
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color
        elif ttype == 'bold':
            run = paragraph.add_run(token[1])
            run.bold = True
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color
        elif ttype == 'italic':
            run = paragraph.add_run(token[1])
            run.italic = True
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color
        elif ttype == 'code':
            run = paragraph.add_run(token[1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 30, 80) # berry / pink code color
        elif ttype == 'link':
            run = paragraph.add_run(token[1])
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = BLUE_RGB
            run.underline = True

def add_code_block(doc, code_text):
    """Renders a code block inside a styled container."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, BG_CODE_HEX)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left accent border
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="16" w:space="0" w:color="{BLUE_HEX}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.rstrip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    # Space after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def add_blockquote(doc, text):
    """Renders a styled blockquote."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EFF6FF") # light blue
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="18" w:space="0" w:color="{NAVY_HEX}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, text, base_font="Calibri", base_size=Pt(10), base_color=RGBColor(30, 58, 138))
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_image_with_caption(doc, image_rel_path, alt_text=""):
    """Adds an image centered with a styled caption."""
    image_full_path = os.path.join(WORKSPACE_DIR, image_rel_path)
    if not os.path.exists(image_full_path):
        print(f"Warning: Image not found at {image_full_path}")
        return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(image_full_path, width=Inches(6.0))
    
    if alt_text:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(alt_text)
        run_cap.italic = True
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9)
        run_cap.font.color.rgb = MUTED_TEXT_RGB

def add_markdown_table(doc, table_lines):
    """Renders a parsed markdown table."""
    # Filter rows and split cells
    rows_data = []
    alignments = []
    
    for line in table_lines:
        line_str = line.strip()
        if not line_str.startswith("|"):
            continue
        cells = [c.strip() for c in line_str.split("|")[1:-1]]
        
        # Check if separator row
        if all(re.match(r"^:?-+:?$", c) for c in cells if c):
            for c in cells:
                if c.startswith(":") and c.endswith(":"):
                    alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
                elif c.endswith(":"):
                    alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
                else:
                    alignments.append(WD_ALIGN_PARAGRAPH.LEFT)
            continue
        rows_data.append(cells)
    
    if not rows_data:
        return
    
    num_cols = max(len(r) for r in rows_data)
    # Normalize rows
    for r in rows_data:
        while len(r) < num_cols:
            r.append("")
            
    while len(alignments) < num_cols:
        alignments.append(WD_ALIGN_PARAGRAPH.LEFT)
        
    doc_table = doc.add_table(rows=len(rows_data), cols=num_cols)
    doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(doc_table)
    
    for row_idx, row in enumerate(rows_data):
        doc_row = doc_table.rows[row_idx]
        
        # CantSplit so row doesn't break across pages
        trPr = doc_row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        is_header = (row_idx == 0)
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for col_idx, cell_text in enumerate(row):
            cell = doc_row.cells[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            
            if is_header:
                set_cell_background(cell, NAVY_HEX)
            elif row_idx % 2 == 1:
                set_cell_background(cell, BG_ALT_ROW_HEX)
            else:
                set_cell_background(cell, "FFFFFF")
                
            p = cell.paragraphs[0]
            p.alignment = alignments[col_idx]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            
            if is_header:
                run = p.add_run(cell_text)
                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                add_inline_runs(p, cell_text, base_font="Calibri", base_size=Pt(9.5), base_color=DARK_TEXT_RGB)
                
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def generate_docx():
    print("Reading README.md...")
    with open(README_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()
    
    # Configure 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Add running header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("AeroFS: Operating System Storage Subsystem Report")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = MUTED_TEXT_RGB
        
        # Add running footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("BITE303L — Operating Systems | VIT University 2025–26")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_TEXT_RGB

    lines = md_text.splitlines()
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Code block handling
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Empty line
        if not stripped:
            i += 1
            continue
            
        # Horizontal rule
        if re.match(r"^---+$", stripped):
            # Add a subtle divider paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="{BORDER_HEX}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue
            
        # Table block detection
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, table_lines)
            continue
            
        # Image detection: ![alt](url)
        img_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if img_match:
            alt = img_match.group(1)
            src = img_match.group(2)
            add_image_with_caption(doc, src, alt)
            i += 1
            continue
            
        # Blockquote: > text
        if stripped.startswith(">"):
            quote_text = stripped.lstrip("> ").strip()
            add_blockquote(doc, quote_text)
            i += 1
            continue
            
        # Headings
        if stripped.startswith("#"):
            match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if match:
                level = len(match.group(1))
                h_text = match.group(2)
                p = doc.add_paragraph()
                
                if level == 1:
                    p.paragraph_format.space_before = Pt(16)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(h_text)
                    run.bold = True
                    run.font.name = "Calibri"
                    run.font.size = Pt(22)
                    run.font.color.rgb = NAVY_RGB
                elif level == 2:
                    p.paragraph_format.space_before = Pt(14)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(h_text)
                    run.bold = True
                    run.font.name = "Calibri"
                    run.font.size = Pt(15)
                    run.font.color.rgb = BLUE_RGB
                elif level == 3:
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(h_text)
                    run.bold = True
                    run.font.name = "Calibri"
                    run.font.size = Pt(12)
                    run.font.color.rgb = RGBColor(30, 41, 59)
                else: # level 4 or higher
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(h_text)
                    run.bold = True
                    run.italic = True
                    run.font.name = "Calibri"
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(71, 85, 105)
                i += 1
                continue

        # Bullet list: - or *
        if stripped.startswith("- ") or stripped.startswith("* "):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, item_text, base_font="Calibri", base_size=Pt(10.5), base_color=DARK_TEXT_RGB)
            i += 1
            continue
            
        # Numbered list: 1. or 2.
        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num_match:
            item_text = num_match.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, item_text, base_font="Calibri", base_size=Pt(10.5), base_color=DARK_TEXT_RGB)
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_runs(p, stripped, base_font="Calibri", base_size=Pt(10.5), base_color=DARK_TEXT_RGB)
        i += 1

    print(f"Saving to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print(f"Saving copy to {OUTPUT_DOCX_ALT}...")
    doc.save(OUTPUT_DOCX_ALT)
    print("Done! Word document generated successfully.")

if __name__ == "__main__":
    generate_docx()
